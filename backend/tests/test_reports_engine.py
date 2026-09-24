"""Comprehensive multi-domain tests for DataScope Unified Reports Module.

Tests:
1. E-commerce dataset (sales, profit, discount, category, region, date)
2. Dataset without profit (handles non-financial gracefully)
3. Dataset without date (handles non-temporal gracefully)
4. Dataset without competitor benchmark (honest missing data notification)
5. Dataset with missing values & duplicate rows
6. Alternative domain dataset (IPL Cricket match analytics)
7. DOCX Word report generation (validates binary stream & XML structure)
8. PDF ReportLab report generation (validates binary stream & %PDF header)
"""
from __future__ import annotations

import io
import unittest
import pandas as pd
from docx import Document

from app.services import dataset_service, dataset_store
from app.services.report_engine import generate_comprehensive_report
from app.services.docx_report_generator import generate_docx_report
from app.services.pdf_report_generator import generate_pdf_report


class TestReportsEngine(unittest.TestCase):
    """Test suite for DataScope comprehensive reporting and exports."""

    def setUp(self):
        dataset_store.clear_store()

    def test_ecommerce_full_report(self):
        """Test complete report on an e-commerce dataset with financial and time series data."""
        df = pd.DataFrame({
            "Order_Date": pd.date_range("2024-01-01", periods=120, freq="D").strftime("%Y-%m-%d"),
            "Category": ["Technology", "Furniture", "Office Supplies", "Apparel"] * 30,
            "Region": ["North", "South", "East", "West"] * 30,
            "Sales": [1200.0, 450.0, 150.0, 80.0] * 30,
            "Profit": [320.0, -45.0, 30.0, 15.0] * 30,
            "Discount": [0.05, 0.25, 0.10, 0.0] * 30,
            "Quantity": [3, 2, 5, 1] * 30,
        })
        dataset_id = dataset_store.save_dataset("superstore_sales.csv", "csv", df)
        report = generate_comprehensive_report(dataset_id)

        # 1. Check metadata
        self.assertEqual(report.metadata.row_count, 120)
        self.assertEqual(report.metadata.column_count, 7)
        self.assertEqual(report.metadata.dataset_name, "superstore_sales.csv")
        self.assertTrue(report.metadata.available_sections_count >= 8)

        # 2. Check Executive Summary
        self.assertTrue(len(report.executive_summary.key_findings) > 0)
        self.assertIn("superstore_sales.csv", report.executive_summary.dataset_name)

        # 3. Check Business Performance & P&L
        self.assertTrue(report.business_performance.is_available)
        self.assertIsNotNone(report.business_performance.total_sales_revenue)
        self.assertTrue(report.profit_loss.is_available)
        self.assertTrue(len(report.profit_loss.loss_making_segments) > 0)

        # 4. Check Trends & Forecasting
        self.assertTrue(report.trends_intelligence.is_available)
        self.assertIsNotNone(report.trends_intelligence.period_comparison)
        self.assertTrue(report.forecasting.is_available)
        self.assertTrue(len(report.forecasting.forecast_points) > 0)

        # 5. Check Recommendations & Action Plan
        self.assertTrue(report.recommendations.is_available)
        self.assertTrue(len(report.recommendations.recommendations_list) > 0)
        self.assertTrue(report.corrective_action_plan.is_available)
        self.assertTrue(len(report.corrective_action_plan.action_items) > 0)
        
        # Verify Action Plan contains priority, problem, recommended_action, owner_team, metric_to_track
        first_action = report.corrective_action_plan.action_items[0]
        self.assertIn(first_action.priority, ["Critical", "High", "Medium", "Low"])
        self.assertTrue(len(first_action.problem) > 0)
        self.assertTrue(len(first_action.recommended_action) > 0)
        self.assertTrue(len(first_action.owner_team) > 0)

        # 6. Check Competition (no external benchmark was provided)
        self.assertFalse(report.market_competition.has_external_benchmark)
        self.assertTrue(len(report.market_competition.required_market_fields) > 0)

    def test_dataset_without_profit(self):
        """Test dataset without profit metrics handles P&L gracefully."""
        df = pd.DataFrame({
            "Date": pd.date_range("2024-01-01", periods=60, freq="D").strftime("%Y-%m-%d"),
            "Product": ["A", "B", "C"] * 20,
            "Units_Sold": [10, 20, 15] * 20,
        })
        dataset_id = dataset_store.save_dataset("units_data.csv", "csv", df)
        report = generate_comprehensive_report(dataset_id)

        # Profit & Loss should report unavailable without throwing errors
        self.assertFalse(report.profit_loss.is_available)
        self.assertIsNotNone(report.profit_loss.unavailable_reason)
        # Overview, Trends, Forecast, Recs, Conclusion should still succeed
        self.assertEqual(report.dataset_overview.row_count, 60)
        self.assertTrue(report.trends_intelligence.is_available)
        self.assertTrue(len(report.data_driven_conclusion.main_findings) > 0)

    def test_dataset_without_date(self):
        """Test dataset without date column degrades trends/forecasting gracefully."""
        df = pd.DataFrame({
            "Department": ["Sales", "Engineering", "HR", "Marketing"] * 15,
            "Salary": [80000, 120000, 65000, 75000] * 15,
            "Tenure_Months": [36, 60, 24, 48] * 15,
            "Performance_Rating": [4.2, 4.8, 3.9, 4.1] * 15,
        })
        dataset_id = dataset_store.save_dataset("hr_workforce.csv", "csv", df)
        report = generate_comprehensive_report(dataset_id)

        self.assertFalse(report.dataset_overview.has_date_dimension)
        self.assertFalse(report.trends_intelligence.is_available)
        self.assertFalse(report.forecasting.is_available)
        self.assertIsNotNone(report.trends_intelligence.unavailable_reason)
        self.assertIsNotNone(report.forecasting.unavailable_reason)
        # Business performance and recommendations should still function
        self.assertTrue(report.business_performance.is_available)
        self.assertTrue(len(report.recommendations.recommendations_list) > 0)

    def test_dataset_with_missing_values_and_duplicates(self):
        """Test dataset with missing values and duplicates reflects in quality report."""
        df = pd.DataFrame({
            "ID": [1, 2, 2, 3, None, 5],
            "Score": [90.0, 85.0, 85.0, None, 70.0, 95.0],
            "Subject": ["Math", "Physics", "Physics", "Chemistry", "Biology", None],
        })
        dataset_id = dataset_store.save_dataset("grades.csv", "csv", df)
        report = generate_comprehensive_report(dataset_id)

        self.assertEqual(report.dataset_overview.row_count, 6)
        self.assertTrue(report.dataset_overview.missing_cells_count > 0)
        self.assertTrue(report.dataset_overview.duplicate_rows_count > 0)
        self.assertTrue(len(report.dataset_overview.important_limitations) > 0)

    def test_cricket_sports_domain_report(self):
        """Test alternative non-ecommerce domain (IPL sports analytics)."""
        df = pd.DataFrame({
            "id": list(range(1, 41)),
            "season": [2024] * 40,
            "team1": ["Chennai Super Kings", "Mumbai Indians", "Royal Challengers Bengaluru", "Kolkata Knight Riders"] * 10,
            "team2": ["Mumbai Indians", "Kolkata Knight Riders", "Chennai Super Kings", "Delhi Capitals"] * 10,
            "toss_winner": ["Chennai Super Kings", "Mumbai Indians", "Royal Challengers Bengaluru", "Kolkata Knight Riders"] * 10,
            "toss_decision": ["bat", "field", "field", "bat"] * 10,
            "winner": ["Chennai Super Kings", "Mumbai Indians", "Chennai Super Kings", "Kolkata Knight Riders"] * 10,
            "result_margin": [20, 5, 8, 14] * 10,
            "player_of_match": ["MS Dhoni", "Rohit Sharma", "Virat Kohli", "Andre Russell"] * 10,
            "venue": ["Wankhede Stadium", "MA Chidambaram Stadium", "M Chinnaswamy Stadium", "Eden Gardens"] * 10,
        })
        dataset_id = dataset_store.save_dataset("ipl_matches_2024.csv", "csv", df)
        report = generate_comprehensive_report(dataset_id)

        self.assertEqual(report.metadata.domain_id, "sports")
        self.assertIn("Cricket", report.metadata.domain_name)
        self.assertTrue(len(report.executive_summary.key_findings) > 0)

    def test_docx_export_generation(self):
        """Test DOCX document generation produces a valid Word document."""
        df = pd.DataFrame({
            "Date": pd.date_range("2024-01-01", periods=40, freq="D").strftime("%Y-%m-%d"),
            "Category": ["Hardware", "Software"] * 20,
            "Revenue": [5000.0, 12000.0] * 20,
            "Profit": [1200.0, 3500.0] * 20,
        })
        dataset_id = dataset_store.save_dataset("tech_sales.csv", "csv", df)
        buf, filename = dataset_service.export_report_docx(dataset_id)

        self.assertTrue(filename.startswith("DataScope_Report_tech_sales_"))
        self.assertTrue(filename.endswith(".docx"))
        
        # Verify valid docx by parsing it with python-docx
        buf.seek(0)
        doc = Document(buf)
        self.assertTrue(len(doc.paragraphs) > 5)
        self.assertTrue(len(doc.tables) >= 2)

    def test_pdf_export_generation(self):
        """Test PDF document generation produces a valid PDF file."""
        df = pd.DataFrame({
            "Date": pd.date_range("2024-01-01", periods=40, freq="D").strftime("%Y-%m-%d"),
            "Category": ["Hardware", "Software"] * 20,
            "Revenue": [5000.0, 12000.0] * 20,
            "Profit": [1200.0, 3500.0] * 20,
        })
        dataset_id = dataset_store.save_dataset("tech_sales.csv", "csv", df)
        buf, filename = dataset_service.export_report_pdf(dataset_id)

        self.assertTrue(filename.startswith("DataScope_Report_tech_sales_"))
        self.assertTrue(filename.endswith(".pdf"))
        
        # Verify PDF magic bytes '%PDF-'
        buf.seek(0)
        content = buf.getvalue()
        self.assertTrue(content.startswith(b"%PDF-"))
        self.assertTrue(len(content) > 1000)


if __name__ == "__main__":
    unittest.main()
